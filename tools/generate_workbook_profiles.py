"""Read the agreed workbook (never edit it) and emit schema-v2 hardware targets.

Requires openpyxl. Formula cells must have cached results from Excel.
The planning object is documentation, not a claim of implemented RTL features.
"""
import argparse
import copy
import hashlib
import json
import math
from pathlib import Path

import openpyxl


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output-dir", type=Path,
                        default=Path(__file__).resolve().parents[1] / "config")
    args = parser.parse_args()
    book = openpyxl.load_workbook(args.workbook, read_only=True, data_only=True)
    sheet = book["三档参数对照"]
    baseline = json.loads((Path(__file__).resolve().parents[1] /
                           "config/ftlpu-lpu32.json").read_text(encoding="utf-8"))
    for col, name in ((3, "ftlpu-lpu64-cloud"), (4, "ftlpu-lpu32-edge"),
                      (5, "ftlpu-lpu16-decode")):
        def value(row):
            result = sheet.cell(row, col).value
            if result is None:
                raise ValueError(f"Missing cached workbook value: {sheet.cell(row, col).coordinate}")
            return result
        config = copy.deepcopy(baseline)
        config["target"]["name"] = name
        config["topology"].update(hemispheres=value(12), tiles_per_slice=value(13),
                                  lanes_per_tile=value(14))
        config["mem"].update(slices_per_hemisphere=value(29), banks_per_slice=value(31),
                             rows_per_bank=value(32), bytes_per_lane=value(33))
        config["stream_topology"]["bindings"]["mem"]["slices_per_group"] = value(30)
        config["sr"]["registers_per_lane"] = value(43)
        config["c2c"]["beat_bytes"] = value(44)
        config["mxm"].update(instances_per_hemisphere=value(16), accum_contexts=value(21))
        config["vxm"]["alus"] = value(25)
        config["throughput"].update(mxms_per_hemisphere=value(16), mxm_weight_buffers=value(20))
        # Layout seeds, not a general tensor lifetime allocator. Keeping legacy
        # addresses 7936/1600 would already be out of bounds on the 1024-row target.
        config["memory"]["attention_mask_base_row"] = value(32) - 256
        config["memory"]["matmul_result_base_row"] = min(1600, value(32) // 2)
        config["external_memory"].update(
            lpu_clock_mhz=value(26),
            ddr_peak_bandwidth_mbytes_per_second=round(min(value(54), value(47)) * 1000),
            ddr_scheduling_efficiency_percent=round(value(55) * 100),
            ddr_read_latency_cycles=math.ceil(value(57) * value(26) / 1000),
            ddr_write_latency_cycles=math.ceil(value(58) * value(26) / 1000),
            ddr_read_latency_jitter_cycles=0, ddr_write_latency_jitter_cycles=0,
            ddr_request_queue_depth=value(60))
        assert value(15) == value(13) * value(14)
        assert value(35) == value(12)*value(29)*value(31)*value(32)*value(34)/1048576
        config["planning"] = {
            "source_workbook": args.workbook.name,
            "source_sha256": hashlib.sha256(args.workbook.read_bytes()).hexdigest(),
            "source_sheet": sheet.title,
            "source_column": sheet.cell(5, col).column_letter,
            "status": "experimental target; model-level validation required",
            "workbook_rows": {str(r): {"parameter": sheet.cell(r, 1).value,
                                       "value": value(r)} for r in range(7, 61)
                              if sheet.cell(r, col).value is not None},
            "unimplemented_requirements": [
                "RTL/PPA/FPGA acceptance and timing closure",
                "iMEM capacity from program peak and FPGA segmented loading",
                "external-memory capacity/type enforcement (budget, efficiency and latency are modeled)"],
            "layout_note": "Mask starts at depth-256; result seed=min(1600,depth/2). These seeds do not prove tensor lifetime fit.",
            "timing_note": "No jitter specified in workbook; jitter=0. Effective peak=min(per-LPU budget,RX entrance); then apply efficiency, matching workbook row 56. Raw interface peak and budget remain in workbook_rows.",
        }
        args.output_dir.mkdir(parents=True, exist_ok=True)
        dest = args.output_dir / (name + ".json")
        dest.write_text(json.dumps(config, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        print(dest)
    book.close()


if __name__ == "__main__":
    main()
